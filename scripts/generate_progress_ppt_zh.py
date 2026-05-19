from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

root = Path(__file__).resolve().parents[1]
out_path = root / "docs" / "项目进度汇报_2026-05-19_中文版.pptx"

prs = Presentation()


def add_title_slide(title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(title: str, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    first = True
    for line in bullets:
        if first:
            tf.text = line
            first = False
        else:
            p = tf.add_paragraph()
            p.text = line
        tf.paragraphs[-1].font.size = Pt(20)


def add_image_slide(title: str, image_rel: str, caption: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    image_path = root / image_rel
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.6), Inches(1.3), width=Inches(12.1), height=Inches(5.1))
    cap = slide.shapes.add_textbox(Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5))
    cap.text_frame.text = caption
    cap.text_frame.paragraphs[0].font.size = Pt(16)


add_title_slide("水介质测试系统项目进度汇报", "日期：2026-05-19")

add_bullets_slide(
    "项目目标与当前状态",
    [
        "目标：落地 Terminal-Station 分布式架构，支持多操作台并行",
        "当前：主框架已完成，联调修复持续推进",
        "进展：关键控制链路与监控链路已可用",
    ],
)

add_bullets_slide(
    "近期关键修复（已完成）",
    [
        "1号台继电器命名与图元显示名称统一",
        "电磁阀3/4 地址映射与图元绑定关系修正",
        "自检补齐电磁阀4静态检查与三阀联动步骤",
        "步骤3/4 压力变化显示与绝对值判定逻辑修复",
    ],
)

add_image_slide("系统界面与架构参考", "docs/UI-design/industrial_control_system.png", "图：工业控制界面参考")
add_image_slide("现场接线与部署关系", "docs/接线图.png", "图：接线图（部署与联调用）")
add_image_slide("设备侧实物验证（流量计1）", "docs/流量计1.jpg", "图：现场设备实拍")
add_image_slide("设备侧实物验证（流量计2）", "docs/流量计2.jpg", "图：现场设备实拍")

add_bullets_slide(
    "下阶段计划（2-6周）",
    [
        "完成 Dashboard、PLC连接、Station管理页面可用版本",
        "打通实时数据链路并完善动作日志可追溯",
        "推进稳定性、断连重连与长稳运行验证",
        "形成试运行与验收资料包",
    ],
)

add_bullets_slide(
    "结论",
    [
        "项目整体按计划推进，架构方向正确",
        "关键联调问题已建立稳定闭环机制",
        "下一阶段聚焦功能完整度与运行稳定性",
    ],
)

out_path.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out_path))
print(f"PPTX_CREATED:{out_path}")
