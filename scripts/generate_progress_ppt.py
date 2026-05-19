from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

root = Path(__file__).resolve().parents[1]
out_path = root / "docs" / "project_progress_2026-05-19.pptx"

prs = Presentation()


def add_title_slide(title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(title: str, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    first = True
    for line in bullets:
        if first:
            tf.text = line
            first = False
        else:
            p = tf.add_paragraph()
            p.text = line
        tf.paragraphs[-1].font.size = Pt(20)


def add_image_slide(title: str, image_rel: str, caption: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    image_path = root / image_rel
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.6), Inches(1.3), width=Inches(12.1), height=Inches(5.1))
    cap = slide.shapes.add_textbox(Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5))
    cap.text_frame.text = caption
    cap.text_frame.paragraphs[0].font.size = Pt(16)


add_title_slide("Water Medium Test System - Progress Report", "Date: 2026-05-19")

add_bullets_slide(
    "Project Goal and Status",
    [
        "Goal: deliver Terminal-Station distributed architecture for multi-station operation",
        "Status: main framework completed and integration fixes are ongoing",
        "Progress: critical control and monitoring chains are available",
    ],
)

add_bullets_slide(
    "Recent Completed Fixes",
    [
        "Unified relay labels with displayed glyph names on Station 1",
        "Corrected Valve 3/4 PLC mapping and relay-glyph binding",
        "Added missing Valve 4 static check and 3-valve linkage test",
        "Fixed pressure change wording and absolute-delta judgment logic",
    ],
)

add_image_slide("UI and Architecture Reference", "docs/UI-design/industrial_control_system.png", "Figure: industrial control UI reference")
add_image_slide("Wiring and Deployment", "docs/接线图.png", "Figure: wiring diagram for deployment and commissioning")
add_image_slide("Field Device Verification - Flowmeter 1", "docs/流量计1.jpg", "Figure: on-site device image")
add_image_slide("Field Device Verification - Flowmeter 2", "docs/流量计2.jpg", "Figure: on-site device image")

add_bullets_slide(
    "Next Stage Plan (2-6 weeks)",
    [
        "Complete usable versions of Dashboard, PLC Connection, and Station Manager pages",
        "Finish real-time data chain integration and action logging traceability",
        "Run stability tests including reconnect and long-running verification",
        "Prepare trial-run and acceptance deliverables",
    ],
)

add_bullets_slide(
    "Conclusion",
    [
        "Project is moving forward with the planned architecture direction",
        "Key integration issues are being closed in a stable feedback loop",
        "Next focus: feature completeness and operational stability",
    ],
)

out_path.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out_path))
print(f"PPTX_CREATED:{out_path}")
